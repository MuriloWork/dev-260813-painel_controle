import openpyxl
import os
import json

def parse_sheet_to_json(file_path, sheet_name):
    if not os.path.exists(file_path):
        print(f"File not found: {file_path}")
        return

    # Carregar workbook sem data_only para ler fórmulas
    wb = openpyxl.load_workbook(file_path, data_only=False)
    
    if sheet_name not in wb.sheetnames:
        print(f"Sheet {sheet_name} not found. Available sheets: {wb.sheetnames}")
        return

    sheet = wb[sheet_name]
    
    # Mapear Letra da Coluna -> Cabeçalho (todas as colunas até a última preenchida)
    from openpyxl.utils import get_column_letter
    col_mapping = {}
    for col_idx in range(1, sheet.max_column + 1):
        letter = get_column_letter(col_idx)
        val = sheet.cell(row=1, column=col_idx).value
        col_mapping[letter] = val if val else f"[Col {letter}]"
    
    # Valor de P1 para a regra de Insert/Delete
    p1_value = sheet['P1'].value
    print(f"Value of P1: {p1_value}")
    
    data = []
    # Assumindo que a primeira linha é o cabeçalho
    headers = [cell.value for cell in sheet[1]]
    
    for row in sheet.iter_rows(min_row=2, max_row=10):
        row_data = {}
        for col_idx in range(1, sheet.max_column + 1):
            letter = get_column_letter(col_idx)
            header = col_mapping[letter]
            cell = row[col_idx-1]
            row_data[f"{letter} ({header})"] = cell.value
        data.append(row_data)
            
    return {"col_mapping": col_mapping, "rows": data}

if __name__ == "__main__":
    # ... (mesmo código de caminhos)
    base_path = r"C:\Users\muril\OneDrive\01 mycloud\01 sistMu\10.01 scripts\01_root\main\260402_crm_cdd"
    xls_path = os.path.join(base_path, "sprints", "260415_srd_analise_metodos_sync_step_3.xlsx")
    sheet = "new_sync"
    
    result_obj = parse_sheet_to_json(xls_path, sheet)
    
    if result_obj:
        # Mostrar as primeiras linhas completas para análise
        print("\nFirst 3 rows data:")
        for i, row in enumerate(result_obj['rows'][:3]):
             print(f"\nRow {i+2}:")
             for k, v in row.items():
                 print(f"  {k}: {v}")

